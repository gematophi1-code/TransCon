from pathlib import Path

from docx import Document
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

_FONT_PATH = Path(__file__).parent / "fonts" / "DejaVuSans.ttf"
_FONT_NAME = "DejaVuSans"
_PAGE_WIDTH, _PAGE_HEIGHT = A4
_MARGIN = 40
_FONT_SIZE = 12
_LINE_HEIGHT = 16

pdfmetrics.registerFont(TTFont(_FONT_NAME, str(_FONT_PATH)))


def _lines_to_pdf(lines: list[str], output_path: Path) -> Path:
    pdf = canvas.Canvas(str(output_path), pagesize=A4)
    pdf.setFont(_FONT_NAME, _FONT_SIZE)

    y = _PAGE_HEIGHT - _MARGIN
    for line in lines:
        if y < _MARGIN:
            pdf.showPage()
            pdf.setFont(_FONT_NAME, _FONT_SIZE)
            y = _PAGE_HEIGHT - _MARGIN
        pdf.drawString(_MARGIN, y, line)
        y -= _LINE_HEIGHT

    pdf.save()
    return output_path


def txt_to_pdf(input_path: Path) -> Path:
    output_path = input_path.with_suffix(".pdf")
    text = input_path.read_text(encoding="utf-8")
    return _lines_to_pdf(text.splitlines() or [""], output_path)


def pdf_to_txt(input_path: Path) -> Path:
    output_path = input_path.with_suffix(".txt")
    reader = PdfReader(str(input_path))
    text_parts = [page.extract_text() or "" for page in reader.pages]
    output_path.write_text("\n".join(text_parts), encoding="utf-8")
    return output_path


def docx_to_txt(input_path: Path) -> Path:
    output_path = input_path.with_suffix(".txt")
    document = Document(input_path)
    lines = [paragraph.text for paragraph in document.paragraphs]
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path


def txt_to_docx(input_path: Path) -> Path:
    output_path = input_path.with_suffix(".docx")
    text = input_path.read_text(encoding="utf-8")

    document = Document()
    for line in text.splitlines() or [""]:
        document.add_paragraph(line)
    document.save(output_path)
    return output_path


def docx_to_pdf(input_path: Path) -> Path:
    output_path = input_path.with_suffix(".pdf")
    document = Document(input_path)
    lines = [paragraph.text for paragraph in document.paragraphs] or [""]
    return _lines_to_pdf(lines, output_path)


def pdf_to_docx(input_path: Path) -> Path:
    output_path = input_path.with_suffix(".docx")
    reader = PdfReader(str(input_path))
    text_parts = [page.extract_text() or "" for page in reader.pages]

    document = Document()
    for line in "\n".join(text_parts).splitlines():
        document.add_paragraph(line)
    document.save(output_path)
    return output_path