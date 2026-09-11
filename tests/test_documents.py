from pathlib import Path

from docx import Document
from pypdf import PdfReader

from src.converter.documents import (
    docx_to_pdf,
    docx_to_txt,
    pdf_to_docx,
    pdf_to_txt,
    txt_to_docx,
    txt_to_pdf,
)


def test_txt_to_pdf_empty_file(tmp_path: Path) -> None:
    # Arrange
    input_path = tmp_path / "empty.txt"
    input_path.write_text("", encoding="utf-8")

    # Act
    output_path = txt_to_pdf(input_path)

    # Assert
    assert output_path.exists()
    reader = PdfReader(str(output_path))
    assert len(reader.pages) == 1


def test_pdf_to_txt_extracts_cyrillic_text(tmp_path: Path) -> None:
    # Arrange
    txt_path = tmp_path / "note.txt"
    txt_path.write_text("Привет, мир!", encoding="utf-8")
    pdf_path = txt_to_pdf(txt_path)

    # Act
    result_path = pdf_to_txt(pdf_path)

    # Assert
    extracted = result_path.read_text(encoding="utf-8")
    assert "Привет" in extracted


def test_txt_to_docx_preserves_lines(tmp_path: Path) -> None:
    # Arrange
    input_path = tmp_path / "note.txt"
    input_path.write_text("Привет, мир!\nвторая строка", encoding="utf-8")

    # Act
    output_path = txt_to_docx(input_path)

    # Assert
    assert output_path.exists()
    document = Document(str(output_path))
    paragraphs = [p.text for p in document.paragraphs]
    assert paragraphs == ["Привет, мир!", "вторая строка"]


def test_docx_to_txt_extracts_paragraphs(tmp_path: Path) -> None:
    # Arrange
    txt_path = tmp_path / "note.txt"
    txt_path.write_text("первая строка\nвторая строка", encoding="utf-8")
    docx_path = txt_to_docx(txt_path)

    # Act
    result_path = docx_to_txt(docx_path)

    # Assert
    extracted = result_path.read_text(encoding="utf-8")
    assert extracted == "первая строка\nвторая строка"


def test_docx_to_pdf_produces_pdf(tmp_path: Path) -> None:
    # Arrange
    txt_path = tmp_path / "note.txt"
    txt_path.write_text("Привет, мир!", encoding="utf-8")
    docx_path = txt_to_docx(txt_path)

    # Act
    output_path = docx_to_pdf(docx_path)

    # Assert
    assert output_path.exists()
    reader = PdfReader(str(output_path))
    assert len(reader.pages) == 1


def test_pdf_to_docx_extracts_text(tmp_path: Path) -> None:
    # Arrange
    txt_path = tmp_path / "note.txt"
    txt_path.write_text("Привет, мир!", encoding="utf-8")
    pdf_path = txt_to_pdf(txt_path)

    # Act
    output_path = pdf_to_docx(pdf_path)

    # Assert
    assert output_path.exists()
    document = Document(str(output_path))
    extracted = "\n".join(p.text for p in document.paragraphs)
    assert "Привет" in extracted